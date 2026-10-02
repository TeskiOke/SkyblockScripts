import os
import io
import math
from pypdf import PdfReader
from docx import Document
from docx.shared import Inches
from PIL import Image

def create_collage(image_data_list, cols=4):
    """Stitches multiple image data buffers into a single PIL Image grid."""
    if not image_data_list:
        return None
    
    images = []
    for data in image_data_list:
        try:
            img = Image.open(io.BytesIO(data))
            img.verify() # Verify it's an image
            img = Image.open(io.BytesIO(data)) # Re-open because verify() closes it or moves the pointer
            images.append(img)
        except:
            continue
            
    if not images:
        return None
        
    # Resize tiles to a standard width for a consistent grid and smaller file size
    tile_w = 400
    images_resized = []
    for img in images:
        aspect = img.height / img.width
        new_h = int(tile_w * aspect)
        images_resized.append(img.resize((tile_w, new_h), Image.Resampling.LANCZOS))
    
    max_h = max(img.height for img in images_resized)
    
    rows = math.ceil(len(images_resized) / cols)
    
    collage_w = tile_w * cols
    collage_h = max_h * rows
    
    # Create blank canvas
    collage = Image.new('RGB', (collage_w, collage_h), (255, 255, 255))
    
    for idx, img in enumerate(images_resized):
        r = idx // cols
        c = idx % cols
        # Center in the row-height if needed, but here we just paste
        collage.paste(img, (c * tile_w, r * max_h))

        
    return collage

def main():
    folder_path = 'foto'
    output_docx = 'compiled_photos_fast.docx'
    password = 'Telegram_Ramen'
    
    doc = Document()
    doc.add_heading('Compiled Photos (Optimized Collage)', 0)
    
    total_files = 2090
    processed_count = 0
    
    # We will process in chunks and save occasionally to maintain performance
    for i in range(1, total_files + 1):
        filename = f"File{i}.pdf"
        file_path = os.path.join(folder_path, filename)
        
        if not os.path.exists(file_path):
            continue
            
        if i % 10 == 0:
            print(f"Processing {filename} ({i}/{total_files})...")
        
        try:
            reader = PdfReader(file_path)
            if reader.is_encrypted:
                reader.decrypt(password)
                
            image_buffers = []
            for page in reader.pages:
                for img_obj in page.images:
                    image_buffers.append(img_obj.data)
            
            if not image_buffers:
                continue
                
            # Stitch images into ONE collage
            collage = create_collage(image_buffers, cols=4)
            if collage:
                # Save collage to a BytesIO object to avoid disk I/O
                combined_img_io = io.BytesIO()
                # Use JPEG to reduce Word doc size
                collage.save(combined_img_io, format='JPEG', quality=85)
                combined_img_io.seek(0)
                
                # Add ONE picture to Word
                doc.add_picture(combined_img_io, width=Inches(6))
                doc.add_paragraph(f"File {i}")
                doc.add_paragraph("-" * 10)
            
            processed_count += 1
            
            # Save less frequently to avoid performance hit
            if processed_count % 100 == 0:
                doc.save(output_docx)
                print(f">>> Saved progress up to File {i} to {output_docx}")
                
        except Exception as e:
            print(f"Error {filename}: {e}")
            
    doc.save(output_docx)
    print(f"Finished! Result saved to {output_docx}")

if __name__ == "__main__":
    main()
